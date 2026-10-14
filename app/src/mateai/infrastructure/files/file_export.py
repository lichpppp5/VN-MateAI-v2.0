# SPDX-License-Identifier: Apache-2.0
# Copyright (c) 2026 Dương Thanh Lịch — VN-MateAI. See LICENSE and NOTICE.
"""
core/file_export.py
===================
Xuất bảng dữ liệu ra CSV / XLSX và đặt tên file tải về an toàn.

Tách khỏi core/server.py (RULE-015) để skill (data_source_tools) dùng được mà
không import mateai.interfaces.http.server.
"""
from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from urllib.parse import quote
from typing import Any, Dict, List

#: Byte Order Mark UTF-8. Excel trên Windows mở file .csv không BOM bằng mã
#: ANSI của máy, nên tiếng Việt ra "Ã¡" hay "?" tuỳ phiên bản. Ghi BOM vào là
#: cách rẻ nhất để file đúng mọi máy — hơn là dặn người dùng mở bằng import.
BOM_UTF8 = b"\xef\xbb\xbf"

def _safe_filename(name: str, max_len: int = 60) -> str:
    """
    Rút gọn tên file về ASCII an toàn.

    Header `Content-Disposition` chỉ mang được ASCII; gửi tiếng Việt thẳng vào
    sẽ bị cắt cụt hoặc làm hỏng header, trình duyệt tải về tên rác. Tên gốc có
    dấu vẫn được giữ qua tham số `filename*` — xem `_content_disposition`.
    """
    raw = str(name or "bao-cao")
    # Bỏ dấu trước, rồi bỏ ký tự lạ — thứ tự này giữ lại được chữ cái.
    ascii_only = unicodedata.normalize("NFKD", raw).encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^A-Za-z0-9_\-]+", "-", ascii_only).strip("-")
    return (cleaned or "bao-cao")[:max_len]


def _content_disposition(title: str, stamp: str, ext: str) -> str:
    """
    Dựng header tải file: `filename` ASCII + `filename*` UTF-8 (RFC 5987/6266).

    Cần cả hai: client cũ chỉ đọc `filename` và sẽ thấy tên không dấu; client
    mới đọc `filename*` và hiện đúng tên có dấu cho người dùng Việt.
    """
    ascii_name = f"{_safe_filename(title)}-{stamp}.{ext}"
    try:
        utf8_name = quote(f"{title}-{stamp}.{ext}", safe="")
    except Exception:  # pragma: no cover - title lạ thì rơi về bản ASCII
        utf8_name = ascii_name
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{utf8_name}"


def _cell_value(value: Any) -> Any:
    """
    Chuẩn hoá một ô trước khi ghi ra file.

    Dict/list không ghi thẳng vào Excel được và cũng vô nghĩa với người đọc
    báo cáo — gộp thành JSON một dòng cho dễ nhìn hơn là "[object Object]".
    """
    if value is None or isinstance(value, (int, float, bool)):
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _rows_to_csv(rows: List[Dict[str, Any]], columns: List[str]) -> bytes:
    """CSV có BOM UTF-8, tiêu đề cột tiếng Việt, dòng \r\n theo thông lệ Excel."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(columns)
    for row in rows:
        writer.writerow([_cell_value(row.get(c)) for c in columns])
    return BOM_UTF8 + buf.getvalue().encode("utf-8")


def _rows_to_xlsx(rows: List[Dict[str, Any]], columns: List[str], sheet_title: str) -> bytes:
    """
    XLSX: dòng tiêu đề đóng băng + tự giãn cột theo nội dung.

    Giãn cột theo độ dài thực tế thay vì đặt cứng — báo cáo tài chính có cột
    "diễn giải" rất dài, đặt cứng sẽ khiến mỗi cột phải mở rộng thủ công.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Báo cáo"  # Excel chặn tên sheet > 31 ký tự

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="334155")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(columns)
    for idx, cell in enumerate(ws[1], start=1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append([_cell_value(row.get(c)) for c in columns])

    for idx, col in enumerate(columns, start=1):
        # Cộng thêm 2 ký tự cho padding, trần 60 để một mô tả dài không đẩy
        # cột khỏi màn hình.
        longest = max([len(str(col))] + [
            len(str(ws.cell(row=r, column=idx).value or "")) for r in range(2, min(ws.max_row, 200) + 1)
        ])
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, 60)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
